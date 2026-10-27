"""AJOU-VIS 100문항 평가 — 평가 문항집(eval/question_bank.xlsx)을 실제 RAG 경로로 한·영 모두 돌리고 채점 초안을 만든다

스모크 테스트(eval/smoke/run_smoke.py)와 같은 방식으로 server/app.py를 프로세스 안에서 호출하고 같은 자동 점검을 적용한다.
답변마다 Claude가 문항집의 기대 정답·정답 판정 기준으로 정답(O/X)·근거 제시(O/X/NA) 초안을 매긴다.
초안은 사람이 검토하고, 검토가 끝난 판정만 문항집에 기록한다 (--fill).

실행 (저장소 루트에서, cmd):
  server\\.venv\\Scripts\\python eval\\bank\\run_bank.py                       한·영 100문항씩 + 채점 초안
  server\\.venv\\Scripts\\python eval\\bank\\run_bank.py --lang ko --only 1,2  일부만
  server\\.venv\\Scripts\\python eval\\bank\\run_bank.py --resume eval\\bank\\results\\<결과>.json    멈춘 실행 이어서 (끝난 문항은 건너뜀)
  server\\.venv\\Scripts\\python eval\\bank\\run_bank.py --rejudge eval\\bank\\results\\<결과>.json   채점 초안만 다시 (답변은 그대로)
  server\\.venv\\Scripts\\python eval\\bank\\run_bank.py --fill eval\\bank\\results\\<결과>.json      검토 반영 판정을 문항집에 기록

결과: eval/bank/results/bank_<날짜-시각>_<모델>.json (원자료·채점 초안, 답변마다 저장) + 같은 이름 .md (언어·영역별 요약, 검토할 문항)
크레딧 부족·인증 오류가 나면 그때까지 저장하고 멈춘다 → --resume으로 이어서 돌린다.
검토: 같은 이름 _review.json에 {"<번호>-<ko|en>": {"correct": "O", "evidence": "NA", "memo": "…"}} 형식으로 바꿀 판정만 적는다.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import statistics
import sys
import time
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
SERVER = ROOT / "server"
HERE = Path(__file__).resolve().parent
RESULTS = HERE / "results"
BANK = ROOT / "eval" / "question_bank.xlsx"

sys.path.insert(0, str(ROOT / "eval" / "smoke"))
import run_smoke as smoke  # noqa: E402  (스모크의 서버 호출·단계 기록·자동 점검을 그대로 쓴다)

AREAS = ["학사규정", "행정절차", "외국인학생행정", "캠퍼스시설"]
FIRST_ROW, LAST_ROW = 4, 103  # 문항집 데이터 행
COL = {"no": 1, "area": 2, "type": 3, "ko": 4, "en": 5, "answer": 6, "doc": 7, "ref": 8,
       "criteria": 9, "correct": 10, "evidence": 11, "summary": 12, "note": 13}

JUDGE_SYSTEM = """너는 아주대학교 캠퍼스 안내 음성 비서 '치토'의 답변을 채점하는 평가자다. 평가 문항집의 기대 정답과 정답 판정 기준만으로 판단한다.

채점 규칙 (평가 문항집 안내 시트와 같음)
- 정답 O/X: 일반 문항은 답변의 사실이 기대 정답·판정 기준에 맞으면 O. 함정-무근거 문항은 규정·수치를 지어내지 않고 담당 부서나 포탈 확인으로 안내하면 O.
- 근거 제시 O/X/NA: 답변이 근거 문서·조항·안내 페이지 이름을 밝혔으면 O, 안 밝혔으면 X. 함정-무근거 문항, 범위 밖 요청(구간 D), 근거가 없어 보류한 답은 NA.
- 답변이 영어여도 같은 기준이다. 한국어 문서명을 영어로 옮겨 말한 것도 근거 제시로 본다.
- 판정 기준이 허용하는 표현 차이와 맞는 추가 사실은 문제 삼지 않는다. 기대 정답에 없는 사실을 지어냈거나 수치가 틀리면 X.
- 판정 기준이 애매하거나 답이 부분적으로만 맞으면 confidence를 low로 한다.

JSON 한 줄로만 답한다 (설명·코드블록 없이):
{"correct": "O|X", "evidence": "O|X|NA", "confidence": "high|low", "reason": "판정 이유 한국어 한 문장", "summary": "실제 응답 요약 한국어 한 줄(40자 안팎)"}"""


def load_bank() -> list[dict]:
    import openpyxl

    ws = openpyxl.load_workbook(BANK, data_only=True)["문항집"]
    rows = []
    for r in range(FIRST_ROW, LAST_ROW + 1):
        q = {k: ws.cell(r, c).value for k, c in COL.items()}
        if not q["no"]:
            continue
        m = re.match(r"구간\s*([A-D])", str(q["note"] or ""))
        q["zone"] = m.group(1) if m else ("C" if q["type"] == "함정-무근거" else "A")
        rows.append(q)
    return rows


def to_case(q: dict, lang: str) -> dict:
    """문항 → 스모크 케이스 형식 (자동 점검용). 정답 여부는 채점 초안이 판단한다."""
    return {
        "id": f"{q['no']}-{lang}", "zone": q["zone"], "lang": lang,
        "question": q[lang], "basis": f"{q['doc']} · {q['ref']}", "expected": q["answer"],
        "cites_source": q["zone"] in ("A", "B"),
    }


class StopRun(Exception):
    """크레딧 부족·인증 실패처럼 계속 돌려도 모두 실패할 오류 — 저장하고 멈춘다."""


def is_fatal(error: str) -> bool:
    return any(k in error for k in ("credit balance", "AuthenticationError", "PermissionDeniedError"))


def judge(client, model: str, effort: str | None, q: dict, rec: dict) -> dict:
    hits = "\n".join(f"- {h['source']} | {h['article']}" for h in rec["hits"]) or "(검색 안 함 또는 게이트 차단)"
    user = (
        f"[문항 {q['no']} · {q['area']} · {q['type']} · 구간 {q['zone']}]\n"
        f"질문({rec['lang']}): {rec['question']}\n"
        f"기대 정답·핵심 사실: {q['answer']}\n"
        f"근거 문서: {q['doc']} · {q['ref']}\n"
        f"정답 판정 기준: {q['criteria']}\n\n"
        f"[치토의 답변]\n{rec['answer'] or '(답변 없음: ' + str(rec.get('error')) + ')'}\n\n"
        f"[참고: 이번 턴 검색결과 상위 3개]\n{hits}"
    )
    kwargs = {"model": model, "max_tokens": 300, "system": JUDGE_SYSTEM,
              "messages": [{"role": "user", "content": user}]}
    if effort and not model.startswith(("claude-haiku-4-5", "claude-sonnet-4-5")):
        kwargs["output_config"] = {"effort": effort}
    for attempt in range(2):
        try:
            resp = client.messages.create(**kwargs)
        except Exception as e:  # 한 문항의 채점 실패가 실행 전체를 멈추지 않게 (크레딧·인증 오류만 멈춤)
            if is_fatal(f"{type(e).__name__}: {e}"):
                raise StopRun(str(e)) from e
            return {"correct": "?", "evidence": "?", "confidence": "low",
                    "reason": f"채점 호출 실패: {type(e).__name__}", "summary": "", "tokens": [0, 0]}
        text = next((b.text for b in resp.content if b.type == "text"), "")
        m = re.search(r"\{.*\}", text, re.S)
        try:
            out = json.loads(m.group(0)) if m else None
        except json.JSONDecodeError:
            out = None
        if out and out.get("correct") in ("O", "X") and out.get("evidence") in ("O", "X", "NA"):
            if q["type"] == "함정-무근거":
                out["evidence"] = "NA"  # 안내 시트 규칙: 함정 문항의 근거 제시는 NA
            out["tokens"] = [resp.usage.input_tokens, resp.usage.output_tokens]
            return out
    return {"correct": "?", "evidence": "?", "confidence": "low", "reason": f"채점 응답 해석 실패: {text[:80]}",
            "summary": "", "tokens": [resp.usage.input_tokens, resp.usage.output_tokens]}


# ── 집계 ────────────────────────────────────────────────────────

def final(rec: dict, review: dict) -> dict:
    """채점 초안에 사람 검토를 덮어쓴 최종 판정."""
    j = dict(rec.get("judge") or {})
    j.update(review.get(rec["id"], {}))
    return j


def rates(recs: list[dict], review: dict) -> dict:
    js = [(r, final(r, review)) for r in recs]
    judged = [(r, j) for r, j in js if j.get("correct") in ("O", "X")]
    normal = [(r, j) for r, j in judged if r["type"] == "일반" and j.get("evidence") in ("O", "X")]
    traps = [(r, j) for r, j in judged if r["type"] == "함정-무근거"]
    pct = lambda a, b: f"{a}/{b} ({a / b:.0%})" if b else "-"  # noqa: E731
    answered = [r for r in recs if r["answer"]]
    return {
        "correct": pct(sum(j["correct"] == "O" for _, j in judged), len(judged)),
        "evidence": pct(sum(j["evidence"] == "O" for _, j in normal), len(normal)),
        "trap": pct(sum(j["correct"] == "O" for _, j in traps), len(traps)),
        "gate_block": sum(r["path"] == "gate_block" for r in recs),
        "auto_fail": sum(not r["auto_pass"] for r in recs),
        "chars": round(statistics.mean(len(r["answer"]) for r in answered)) if answered else 0,
        "words": round(statistics.mean(len(r["answer"].split()) for r in answered)) if answered else 0,
        "seconds": round(statistics.median(r["seconds"] for r in recs), 1) if recs else 0,
    }


def judge_usd(model: str, recs: list[dict]) -> float:
    price = smoke.PRICES.get(model, (0, 0))
    tin = sum((r.get("judge") or {}).get("tokens", [0, 0])[0] for r in recs)
    tout = sum((r.get("judge") or {}).get("tokens", [0, 0])[1] for r in recs)
    return (tin * price[0] + tout * price[1]) / 1_000_000


def md_cell(text, limit=70) -> str:
    return smoke.md_cell(str(text or ""), limit)


def write_report(data: dict, out_md: Path, review: dict):
    meta, recs = data["meta"], data["results"]
    langs = [l for l in ("ko", "en") if any(r["lang"] == l for r in recs)]
    by = {(r["no"], r["lang"]): r for r in recs}
    lines = [
        f"# 100문항 평가 결과 — {meta['started']}",
        "",
        f"- 모델: `{meta['model']}` (effort `{meta['effort']}`) · 프롬프트 {meta['prompt_version']} · 코드 {meta['git']}",
        f"- 문항: {meta['count']}개 × {', '.join(langs)} · 게이트 임계값 {meta['threshold']}"
        + ("" if meta.get("complete", True) else " · **실행이 끝나지 않음 (--resume으로 이어서)**"),
        f"- 비용(추정): 답변 ${meta['usd_answer']:.2f} + 채점 초안 ${meta['usd_judge']:.2f}",
        f"- 사람 검토 반영: {len(review)}건" + (f" (`{meta.get('review_file')}`)" if review else " (아직 없음 — 아래 판정은 Claude 초안)"),
        "",
        "정답률 = 정답 O / 채점한 문항 · 근거 제시율 = 일반 문항 중 근거 O / (O+X), NA 제외 · 무근거 회피율 = 함정 문항 정답 O / 함정 문항",
        "",
        "## 언어별 요약",
        "",
        "| 지표 | " + " | ".join({"ko": "한국어", "en": "영어"}[l] for l in langs) + " |",
        "|---|" + "---|" * len(langs),
    ]
    stats = {l: rates([r for r in recs if r["lang"] == l], review) for l in langs}
    for key, label in (("correct", "정답률"), ("evidence", "근거 제시율"), ("trap", "무근거 회피율"),
                       ("gate_block", "게이트 차단(보류 멘트)"), ("auto_fail", "자동 점검 실패"),
                       ("chars", "평균 길이(자)"), ("words", "평균 길이(단어)"), ("seconds", "응답 시간 중앙값(s)")):
        lines.append(f"| {label} | " + " | ".join(str(stats[l][key]) for l in langs) + " |")

    lines += ["", "## 영역별 정답률", "", "| 영역 | " + " | ".join(langs) + " |", "|---|" + "---|" * len(langs)]
    for area in AREAS:
        lines.append(f"| {area} | " + " | ".join(
            rates([r for r in recs if r["lang"] == l and r["area"] == area], review)["correct"] for l in langs) + " |")

    def judged(r):
        j = final(r, review)
        return f"{j.get('correct', '?')}/{j.get('evidence', '?')}"

    if len(langs) == 2:
        split = [n for n in sorted({r["no"] for r in recs})
                 if (n, "ko") in by and (n, "en") in by
                 and final(by[(n, "ko")], review).get("correct") != final(by[(n, "en")], review).get("correct")]
        lines += ["", f"## 한·영 정답 판정이 갈린 문항 ({len(split)}개)", "",
                  "| 번호 | 영역 | 한국어 | 영어 | 영어 검색어 | 영어 검색 1위 |", "|---|---|---|---|---|---|"]
        for n in split:
            ko, en = by[(n, "ko")], by[(n, "en")]
            top = en["hits"][0]["source"] if en["hits"] else "-"
            lines.append(f"| {n} | {ko['area']} | {judged(ko)} | {judged(en)} | {md_cell(en.get('search_query'), 40)} | {md_cell(top, 30)} |")

    check = [r for r in recs if r["id"] not in review and (
        (r.get("judge") or {}).get("confidence") == "low" or (r.get("judge") or {}).get("correct") == "?")]
    lines += ["", f"## 검토가 필요한 초안 ({len(check)}개: 확신 낮음·해석 실패)", "",
              "| ID | 영역 | 판정 | 이유 |", "|---|---|---|---|"]
    for r in check:
        lines.append(f"| {r['id']} | {r['area']} | {judged(r)} | {md_cell(r['judge'].get('reason'), 90)} |")

    lines += ["", "## 문항별", ""]
    for n in sorted({r["no"] for r in recs}):
        first = by.get((n, "ko")) or by.get((n, "en"))
        lines += [f"### {n} · {first['area']} · {first['type']} · 구간 {first['zone']}", "",
                  f"- 기대: {first['expected']}", f"- 판정 기준: {first['criteria']}"]
        for l in langs:
            r = by.get((n, l))
            if not r:
                continue
            j = final(r, review)
            hits = ", ".join(f"{h['source']} {h['score']:.2f}" for h in r["hits"]) or "-"
            lines += ["", f"**{l}** {r['question']} → 판정 {j.get('correct')}/{j.get('evidence')}"
                      + (" (검토 반영)" if r["id"] in review else f" · 확신 {j.get('confidence')}"),
                      f"- 경로 `{r['path']}` · 검색 {hits} · {r['seconds']}s"
                      + (f" · 검색어 {r['search_query']}" if r.get("search_query") and r["search_query"] != r["question"] else ""),
                      "", f"> {r['answer'] or '(답변 없음)'}", "", f"- 판정 이유: {j.get('reason')}"]
            for i in r["issues"]:
                lines.append(f"- ❌ {i}")
        lines.append("")
    out_md.write_text("\n".join(lines), encoding="utf-8")


# ── 문항집 기록 ─────────────────────────────────────────────────

def fill(data: dict, review: dict):
    """검토 반영 판정을 문항집에 기록: 한국어 → 문항집 J·K·L 열, 영어 → '영어 채점' 시트."""
    import openpyxl

    wb = openpyxl.load_workbook(BANK)
    ws = wb["문항집"]
    row_of = {ws.cell(r, COL["no"]).value: r for r in range(FIRST_ROW, LAST_ROW + 1)}
    en_rows = []
    for r in data["results"]:
        j = final(r, review)
        if j.get("correct") not in ("O", "X"):
            continue
        if r["lang"] == "ko":
            row = row_of[r["no"]]
            ws.cell(row, COL["correct"]).value = j["correct"]
            ws.cell(row, COL["evidence"]).value = j["evidence"]
            ws.cell(row, COL["summary"]).value = j.get("summary") or md_cell(r["answer"], 60)
        else:
            en_rows.append((r["no"], r["area"], r["type"], r["question"], j["correct"], j["evidence"],
                            j.get("summary") or md_cell(r["answer"], 60)))
    if en_rows:
        if "영어 채점" in wb.sheetnames:
            del wb["영어 채점"]
        es = wb.create_sheet("영어 채점")
        es.append([f"한·영 A/B 테스트 — 영어 질문 채점 ({data['meta']['started']}, {data['meta']['model']}, "
                   f"프롬프트 {data['meta']['prompt_version']}). 한국어 채점은 문항집 J·K·L 열"])
        es.append(["번호", "영역", "유형", "질문(영어)", "정답(O/X)", "근거·출처 제시(O/X/NA)", "실제 응답 요약"])
        for row in sorted(en_rows):
            es.append(list(row))
        for col, width in zip("ABCDEFG", (6, 14, 12, 60, 10, 18, 50)):
            es.column_dimensions[col].width = width
    wb.save(BANK)
    print(f"문항집 기록: 한국어 {sum(r['lang'] == 'ko' for r in data['results'])}개 → 문항집 J·K·L, 영어 {len(en_rows)}개 → '영어 채점' 시트")


# ── 실행 ─────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser(description="AJOU-VIS 100문항 평가 (한·영)")
    ap.add_argument("--lang", choices=["ko", "en", "both"], default="both")
    ap.add_argument("--only", help="문항 번호, 쉼표 구분 (예: 1,2,49)")
    ap.add_argument("--model", help="LLM_MODEL 덮어쓰기 (답변·채점 모두)")
    ap.add_argument("--no-judge", action="store_true", help="채점 초안 없이 답변만")
    ap.add_argument("--resume", metavar="RESULT_JSON", help="멈춘 실행을 이어서 (끝난 문항은 건너뜀)")
    ap.add_argument("--rejudge", metavar="RESULT_JSON", help="저장된 답변으로 채점 초안만 다시")
    ap.add_argument("--fill", metavar="RESULT_JSON", help="검토 반영 판정을 문항집에 기록 (LLM 호출 없음)")
    args = ap.parse_args()

    if args.model:
        os.environ["LLM_MODEL"] = args.model
    sys.path.insert(0, str(SERVER))
    from dotenv import load_dotenv
    load_dotenv(SERVER / ".env")

    if args.fill:
        path = Path(args.fill)
        data = json.loads(path.read_text(encoding="utf-8"))
        review_path = path.with_name(path.stem + "_review.json")
        review = json.loads(review_path.read_text(encoding="utf-8")) if review_path.exists() else {}
        fill(data, review)
        data["meta"]["review_file"] = review_path.name if review else None
        write_report(data, path.with_suffix(".md"), review)
        return

    import anthropic
    import app
    import gate
    import llm
    import rewriter
    import smalltalk
    from fastapi.testclient import TestClient

    judge_client = anthropic.Anthropic()  # 답변 호출 기록(스모크 tracing)과 섞이지 않게 따로 만든다
    bank = load_bank()
    q_by_no = {q["no"]: q for q in bank}

    if args.rejudge:
        path = Path(args.rejudge)
        data = json.loads(path.read_text(encoding="utf-8"))
        for r in data["results"]:
            r["judge"] = judge(judge_client, llm.MODEL, llm.EFFORT, q_by_no[r["no"]], r)
        data["meta"]["usd_judge"] = judge_usd(llm.MODEL, data["results"])
        path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        write_report(data, path.with_suffix(".md"), {})
        print(f"채점 초안 다시: {path.with_suffix('.md').resolve().relative_to(ROOT)}")
        return

    if type(app.retriever).__name__ == "StubRetriever":
        sys.exit("corpus/chroma_db가 없습니다. 먼저 corpus/scripts/embed_corpus.py로 검색 DB를 만드세요.")
    if args.only:
        wanted = {int(x) for x in args.only.split(",")}
        bank = [q for q in bank if q["no"] in wanted]
    langs = ["ko", "en"] if args.lang == "both" else [args.lang]

    app.retriever.search("휴학", top_k=1)  # 임베딩 모델을 미리 올린다 (스모크와 같음)
    trace = smoke.Trace()
    smoke.install_tracing(app, gate, rewriter, smalltalk, llm, trace, with_llm=True)
    client = TestClient(app.app)

    RESULTS.mkdir(parents=True, exist_ok=True)
    if args.resume:
        out_json = Path(args.resume)
        prev = json.loads(out_json.read_text(encoding="utf-8"))
        started = datetime.strptime(prev["meta"]["started"], "%Y-%m-%d %H:%M")
        # 답변 오류·채점 실패가 남은 문항은 다시 돌린다
        records = [r for r in prev["results"] if not r.get("error")
                   and (args.no_judge or (r.get("judge") or {}).get("correct") in ("O", "X"))]
    else:
        started = datetime.now()
        out_json = RESULTS / f"bank_{started.strftime('%Y%m%d-%H%M')}_{llm.MODEL}.json"
        records = []
    done = {r["id"] for r in records}

    def save(complete: bool) -> dict:
        _, _, usd = smoke.cost_usd(llm.MODEL, records)
        meta = {
            "started": started.strftime("%Y-%m-%d %H:%M"), "model": llm.MODEL, "effort": llm.EFFORT,
            "threshold": gate.SCORE_THRESHOLD, "prompt_version": smoke.prompt_version(), "git": smoke.git_info(),
            "count": len({r["no"] for r in records}), "complete": complete,
            "usd_answer": usd or 0.0, "usd_judge": judge_usd(llm.MODEL, records),
        }
        data = {"meta": meta, "results": records}
        out_json.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        return data

    todo = [(q, lang) for q in bank for lang in langs if f"{q['no']}-{lang}" not in done]
    print(f"100문항 평가 {len(todo)}건 남음 (끝난 {len(done)}건) · {llm.MODEL} (effort {llm.EFFORT})", flush=True)
    for q, lang in todo:
        rec = smoke.run_case(to_case(q, lang), client, trace, app, smalltalk, gate, dry_run=False)
        rec.update(no=q["no"], area=q["area"], type=q["type"], criteria=q["criteria"])
        if rec.get("error") and is_fatal(rec["error"]):
            save(complete=False)
            sys.exit(f"API 오류로 멈춤 ({rec['error'][:120]}) — 크레딧·키를 확인한 뒤 --resume {out_json}")
        try:
            if not args.no_judge:
                rec["judge"] = judge(judge_client, llm.MODEL, llm.EFFORT, q, rec)
        except StopRun as e:
            save(complete=False)
            sys.exit(f"API 오류로 멈춤 ({str(e)[:120]}) — 크레딧·키를 확인한 뒤 --resume {out_json}")
        records.append(rec)
        save(complete=False)  # 답변마다 저장 — 중간에 멈춰도 이어서 돌릴 수 있게
        j = rec.get("judge") or {}
        print(f"  {rec['id']:<7} {rec['path']:<10} top {rec['top_score']:.2f} {rec['seconds']:>5.1f}s "
              f"판정 {j.get('correct', '-')}/{j.get('evidence', '-')} {j.get('confidence', '')}", flush=True)

    order = {(q["no"], lang): i for i, (q, lang) in enumerate((q, l) for q in bank for l in langs)}
    records.sort(key=lambda r: order.get((r["no"], r["lang"]), 0))
    data = save(complete=True)
    write_report(data, out_json.with_suffix(".md"), {})
    print(f"비용(추정) 답변 ${data['meta']['usd_answer']:.2f} + 채점 ${data['meta']['usd_judge']:.2f}")
    print(f"결과: {out_json.with_suffix('.md').resolve().relative_to(ROOT)}")


if __name__ == "__main__":
    main()
